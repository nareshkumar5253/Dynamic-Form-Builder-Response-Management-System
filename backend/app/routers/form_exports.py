
from io import BytesIO
import json

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, selectinload

from app.core.database import get_db
from app.core.deps import require_admin
from app.models.form import Form
from app.models.form_response import FormResponse
from app.models.user import User

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)


router = APIRouter(
    prefix="/forms",
    tags=["Form Exports"],
)


def readable_value(value):
    """Convert dynamic response values to readable text."""
    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)

    return str(value)


def excel_safe_value(value):
    """Prevent response text from being interpreted as an Excel formula."""
    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        value = "'" + value

    return value


def get_form_export_data(form_id: int, db: Session):
    """Load a form and its responses for export."""
    form = db.query(Form).filter(Form.id == form_id).first()

    if form is None:
        raise HTTPException(
            status_code=404,
            detail="Form not found",
        )

    responses = (
        db.query(FormResponse)
        .options(
            selectinload(FormResponse.details),
            selectinload(FormResponse.user),
        )
        .filter(FormResponse.form_id == form_id)
        .order_by(FormResponse.id.asc())
        .all()
    )

    fields = sorted(
        form.fields,
        key=lambda field: (field.field_order, field.id),
    )

    return form, fields, responses


@router.get("/{form_id}/export/responses.xlsx")
def export_form_responses_excel(
    form_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    form, fields, responses = get_form_export_data(form_id, db)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Form Responses"

    headers = [
        "Response ID",
        "Submitted By",
        "Status",
        "Submitted At",
        "Updated At",
    ]
    headers.extend(field.field_label for field in fields)

    worksheet.append([excel_safe_value(header) for header in headers])

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for response in responses:
        details_by_field_id = {
            detail.field_id: detail.value
            for detail in response.details
        }

        submitted_by = (
            response.user.email
            if response.user is not None
            else "Anonymous"
        )

        row = [
            response.id,
            submitted_by,
            response.status,
            response.submitted_at,
            response.updated_at,
        ]

        for field in fields:
            row.append(
                excel_safe_value(
                    details_by_field_id.get(field.id)
                )
            )

        worksheet.append(row)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30

    for column_index, column_cells in enumerate(
        worksheet.columns,
        start=1,
    ):
        max_length = max(
            (
                len(str(cell.value))
                if cell.value is not None
                else 0
            )
            for cell in column_cells
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(max(max_length + 2, 14), 40)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="form_{form.id}_responses.xlsx"'
        },
    )


@router.get("/{form_id}/export/responses.pdf")
def export_form_responses_pdf(
    form_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    form, fields, responses = get_form_export_data(form_id, db)

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"{form.title} - Responses",
        author="Dynamic Form Builder",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        spaceAfter=8,
    )

    normal_style = ParagraphStyle(
        "ExportNormal",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        wordWrap="CJK",
    )

    header_style = ParagraphStyle(
        "ExportHeader",
        parent=normal_style,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=9,
        alignment=TA_CENTER,
    )

    story = [
        Paragraph(
            readable_value(form.title),
            title_style,
        ),
        Paragraph(
            f"Form ID: {form.id} | Total responses: {len(responses)}",
            styles["Normal"],
        ),
        Spacer(1, 5 * mm),
    ]

    headers = [
        "Response ID",
        "Submitted By",
        "Status",
        "Submitted At",
        "Updated At",
    ]
    headers.extend(field.field_label for field in fields)

    table_data = [
        [
            Paragraph(
                readable_value(header).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
                header_style,
            )
            for header in headers
        ]
    ]

    for response in responses:
        details_by_field_id = {
            detail.field_id: detail.value
            for detail in response.details
        }

        submitted_by = (
            response.user.email
            if response.user is not None
            else "Anonymous"
        )

        row_values = [
            response.id,
            submitted_by,
            response.status,
            response.submitted_at.strftime("%Y-%m-%d %H:%M")
            if response.submitted_at else "",
            response.updated_at.strftime("%Y-%m-%d %H:%M")
            if response.updated_at else "",
        ]

        row_values.extend(
            readable_value(details_by_field_id.get(field.id))
            for field in fields
        )

        table_data.append(
            [
                Paragraph(
                    readable_value(value)
                    .replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                    .replace("\n", "<br/>"),
                    normal_style,
                )
                for value in row_values
            ]
        )

    available_width = (
        landscape(A4)[0] - document.leftMargin - document.rightMargin
    )

    column_count = max(len(headers), 1)
    column_width = available_width / column_count
    column_widths = [column_width] * column_count

    response_table = Table(
        table_data,
        colWidths=column_widths,
        repeatRows=1,
        hAlign="LEFT",
    )

    response_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1F4E78"),
                ),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C9D6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F2F6FA")],
                ),
            ]
        )
    )

    story.append(response_table)
    document.build(story)

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="form_{form.id}_responses.pdf"'
        },
    )
